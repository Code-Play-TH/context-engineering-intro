"""
PowerPoint utilities for report generation.
"""
import os
from typing import Dict, Any, List, Optional, Tuple
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
import matplotlib.pyplot as plt
import io
from PIL import Image


class PowerPointStyler:
    """Handles PowerPoint styling and formatting."""
    
    # Default color scheme
    COLORS = {
        'primary': RGBColor(31, 78, 121),      # #1f4e79
        'secondary': RGBColor(112, 173, 71),   # #70ad47
        'accent': RGBColor(255, 192, 0),       # #ffc000
        'danger': RGBColor(197, 90, 90),       # #c55a5a
        'text': RGBColor(68, 68, 68),          # #444444
        'light_gray': RGBColor(242, 242, 242), # #f2f2f2
        'dark_gray': RGBColor(128, 128, 128),  # #808080
    }
    
    # Default fonts
    FONTS = {
        'title': 'Arial',
        'heading': 'Arial',
        'body': 'Arial',
        'caption': 'Arial'
    }
    
    # Font sizes
    FONT_SIZES = {
        'title': Pt(44),
        'subtitle': Pt(24),
        'heading1': Pt(32),
        'heading2': Pt(24),
        'heading3': Pt(18),
        'body': Pt(14),
        'caption': Pt(12),
        'small': Pt(10)
    }
    
    @classmethod
    def apply_text_style(cls, text_frame, style_name: str, text_content: str):
        """Apply text styling to a text frame."""
        paragraph = text_frame.paragraphs[0]
        paragraph.text = text_content
        
        # Set font
        font = paragraph.font
        font.name = cls.FONTS.get('body', 'Arial')
        
        # Set size based on style
        if style_name in cls.FONT_SIZES:
            font.size = cls.FONT_SIZES[style_name]
        else:
            font.size = cls.FONT_SIZES['body']
        
        # Set color based on style
        if style_name in ['title', 'heading1', 'heading2']:
            font.color.rgb = cls.COLORS['primary']
        else:
            font.color.rgb = cls.COLORS['text']
        
        # Set bold for headings
        if style_name in ['title', 'heading1', 'heading2', 'heading3']:
            font.bold = True
        
        # Set alignment
        if style_name == 'title':
            paragraph.alignment = PP_ALIGN.CENTER
        else:
            paragraph.alignment = PP_ALIGN.LEFT
    
    @classmethod
    def get_chart_colors(cls, count: int) -> List[RGBColor]:
        """Get a list of colors for chart series."""
        base_colors = [
            cls.COLORS['primary'],
            cls.COLORS['secondary'],
            cls.COLORS['accent'],
            cls.COLORS['danger'],
            RGBColor(228, 64, 95),   # Instagram pink
            RGBColor(24, 119, 242),  # Facebook blue
            RGBColor(29, 161, 242),  # Twitter blue
            RGBColor(255, 0, 0),     # YouTube red
        ]
        
        # Repeat colors if we need more
        colors = []
        for i in range(count):
            colors.append(base_colors[i % len(base_colors)])
        
        return colors


class PowerPointLayoutManager:
    """Manages PowerPoint slide layouts and positioning."""
    
    # Standard slide dimensions (16:9)
    SLIDE_WIDTH = Inches(13.33)
    SLIDE_HEIGHT = Inches(7.5)
    
    # Standard margins
    MARGIN_LEFT = Inches(0.5)
    MARGIN_RIGHT = Inches(0.5)
    MARGIN_TOP = Inches(0.5)
    MARGIN_BOTTOM = Inches(0.5)
    
    # Content area
    CONTENT_WIDTH = SLIDE_WIDTH - MARGIN_LEFT - MARGIN_RIGHT
    CONTENT_HEIGHT = SLIDE_HEIGHT - MARGIN_TOP - MARGIN_BOTTOM
    
    @classmethod
    def get_position(cls, x: float, y: float, width: float, height: float) -> Tuple[int, int, int, int]:
        """Convert relative positions to absolute positions."""
        abs_x = cls.MARGIN_LEFT + Inches(x)
        abs_y = cls.MARGIN_TOP + Inches(y)
        abs_width = Inches(width)
        abs_height = Inches(height)
        
        return abs_x, abs_y, abs_width, abs_height
    
    @classmethod
    def create_title_layout(cls, slide, title: str, subtitle: str = None):
        """Create a title slide layout."""
        # Title
        title_shape = slide.shapes.add_textbox(
            cls.MARGIN_LEFT,
            Inches(2),
            cls.CONTENT_WIDTH,
            Inches(1.5)
        )
        PowerPointStyler.apply_text_style(title_shape.text_frame, 'title', title)
        
        # Subtitle
        if subtitle:
            subtitle_shape = slide.shapes.add_textbox(
                cls.MARGIN_LEFT,
                Inches(3.8),
                cls.CONTENT_WIDTH,
                Inches(1)
            )
            PowerPointStyler.apply_text_style(subtitle_shape.text_frame, 'subtitle', subtitle)
    
    @classmethod
    def create_content_layout(cls, slide, title: str):
        """Create a content slide layout with title."""
        # Title
        title_shape = slide.shapes.add_textbox(
            cls.MARGIN_LEFT,
            cls.MARGIN_TOP,
            cls.CONTENT_WIDTH,
            Inches(0.8)
        )
        PowerPointStyler.apply_text_style(title_shape.text_frame, 'heading1', title)
        
        return Inches(1.3)  # Return Y position for content


class PowerPointChartGenerator:
    """Generates charts for PowerPoint slides."""
    
    @classmethod
    def add_bar_chart(
        cls, 
        slide, 
        chart_data: Dict[str, Any], 
        x: float, 
        y: float, 
        width: float, 
        height: float
    ):
        """Add a bar chart to the slide."""
        # Prepare chart data
        chart_data_obj = CategoryChartData()
        
        labels = chart_data.get('labels', [])
        datasets = chart_data.get('datasets', [])
        
        # Add categories
        chart_data_obj.categories = labels
        
        # Add series
        for dataset in datasets:
            series_name = dataset.get('label', 'Series')
            series_data = dataset.get('data', [])
            chart_data_obj.add_series(series_name, series_data)
        
        # Create chart
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            abs_x, abs_y, abs_width, abs_height,
            chart_data_obj
        ).chart
        
        # Style the chart
        cls._style_chart(chart, chart_data)
        
        return chart
    
    @classmethod
    def add_line_chart(
        cls, 
        slide, 
        chart_data: Dict[str, Any], 
        x: float, 
        y: float, 
        width: float, 
        height: float
    ):
        """Add a line chart to the slide."""
        # Prepare chart data
        chart_data_obj = CategoryChartData()
        
        labels = chart_data.get('labels', [])
        datasets = chart_data.get('datasets', [])
        
        # Add categories
        chart_data_obj.categories = labels
        
        # Add series
        for dataset in datasets:
            series_name = dataset.get('label', 'Series')
            series_data = dataset.get('data', [])
            chart_data_obj.add_series(series_name, series_data)
        
        # Create chart
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.LINE,
            abs_x, abs_y, abs_width, abs_height,
            chart_data_obj
        ).chart
        
        # Style the chart
        cls._style_chart(chart, chart_data)
        
        return chart
    
    @classmethod
    def add_pie_chart(
        cls, 
        slide, 
        chart_data: Dict[str, Any], 
        x: float, 
        y: float, 
        width: float, 
        height: float
    ):
        """Add a pie chart to the slide."""
        # Prepare chart data
        chart_data_obj = CategoryChartData()
        
        labels = chart_data.get('labels', [])
        datasets = chart_data.get('datasets', [])
        
        if datasets:
            dataset = datasets[0]  # Pie charts only use one dataset
            data = dataset.get('data', [])
            
            # Add categories and data
            chart_data_obj.categories = labels
            chart_data_obj.add_series('Data', data)
        
        # Create chart
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE,
            abs_x, abs_y, abs_width, abs_height,
            chart_data_obj
        ).chart
        
        # Style the chart
        cls._style_chart(chart, chart_data)
        
        return chart
    
    @classmethod
    def add_donut_chart(
        cls, 
        slide, 
        chart_data: Dict[str, Any], 
        x: float, 
        y: float, 
        width: float, 
        height: float
    ):
        """Add a donut chart to the slide."""
        # Prepare chart data
        chart_data_obj = CategoryChartData()
        
        labels = chart_data.get('labels', [])
        datasets = chart_data.get('datasets', [])
        
        if datasets:
            dataset = datasets[0]  # Donut charts only use one dataset
            data = dataset.get('data', [])
            
            # Add categories and data
            chart_data_obj.categories = labels
            chart_data_obj.add_series('Data', data)
        
        # Create chart
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.DOUGHNUT,
            abs_x, abs_y, abs_width, abs_height,
            chart_data_obj
        ).chart
        
        # Style the chart
        cls._style_chart(chart, chart_data)
        
        return chart
    
    @classmethod
    def add_area_chart(
        cls, 
        slide, 
        chart_data: Dict[str, Any], 
        x: float, 
        y: float, 
        width: float, 
        height: float
    ):
        """Add an area chart to the slide."""
        # Prepare chart data
        chart_data_obj = CategoryChartData()
        
        labels = chart_data.get('labels', [])
        datasets = chart_data.get('datasets', [])
        
        # Add categories
        chart_data_obj.categories = labels
        
        # Add series
        for dataset in datasets:
            series_name = dataset.get('label', 'Series')
            series_data = dataset.get('data', [])
            chart_data_obj.add_series(series_name, series_data)
        
        # Create chart
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.AREA,
            abs_x, abs_y, abs_width, abs_height,
            chart_data_obj
        ).chart
        
        # Style the chart
        cls._style_chart(chart, chart_data)
        
        return chart
    
    @classmethod
    def add_stacked_bar_chart(
        cls, 
        slide, 
        chart_data: Dict[str, Any], 
        x: float, 
        y: float, 
        width: float, 
        height: float
    ):
        """Add a stacked bar chart to the slide."""
        # Prepare chart data
        chart_data_obj = CategoryChartData()
        
        labels = chart_data.get('labels', [])
        datasets = chart_data.get('datasets', [])
        
        # Add categories
        chart_data_obj.categories = labels
        
        # Add series
        for dataset in datasets:
            series_name = dataset.get('label', 'Series')
            series_data = dataset.get('data', [])
            chart_data_obj.add_series(series_name, series_data)
        
        # Create chart
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_STACKED,
            abs_x, abs_y, abs_width, abs_height,
            chart_data_obj
        ).chart
        
        # Style the chart
        cls._style_chart(chart, chart_data)
        
        return chart
    
    @classmethod
    def _style_chart(cls, chart, chart_data: Dict[str, Any]):
        """Apply styling to a chart."""
        # Set title
        title = chart_data.get('title', '')
        if title:
            chart.has_title = True
            chart.chart_title.text_frame.text = title
            
            # Style title
            title_font = chart.chart_title.text_frame.paragraphs[0].font
            title_font.size = PowerPointStyler.FONT_SIZES['heading3']
            title_font.color.rgb = PowerPointStyler.COLORS['primary']
            title_font.bold = True
        
        # Set legend
        datasets = chart_data.get('datasets', [])
        if len(datasets) > 1:
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            
            # Style legend
            try:
                legend_font = chart.legend.font
                legend_font.size = PowerPointStyler.FONT_SIZES['caption']
                legend_font.color.rgb = PowerPointStyler.COLORS['text']
            except:
                pass
        else:
            chart.has_legend = False
        
        # Style axes
        try:
            # Category axis (X-axis)
            if hasattr(chart, 'category_axis'):
                cat_axis = chart.category_axis
                cat_axis.has_major_gridlines = False
                cat_axis.tick_labels.font.size = PowerPointStyler.FONT_SIZES['small']
                cat_axis.tick_labels.font.color.rgb = PowerPointStyler.COLORS['text']
            
            # Value axis (Y-axis)
            if hasattr(chart, 'value_axis'):
                val_axis = chart.value_axis
                val_axis.has_major_gridlines = True
                val_axis.major_gridlines.format.line.color.rgb = PowerPointStyler.COLORS['light_gray']
                val_axis.tick_labels.font.size = PowerPointStyler.FONT_SIZES['small']
                val_axis.tick_labels.font.color.rgb = PowerPointStyler.COLORS['text']
        except:
            pass
        
        # Style series colors
        colors = PowerPointStyler.get_chart_colors(len(datasets))
        
        try:
            for i, series in enumerate(chart.series):
                if i < len(colors):
                    # Set series color
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = colors[i]
                    
                    # Set border for better visibility
                    series.format.line.color.rgb = colors[i]
                    series.format.line.width = Pt(1)
        except Exception:
            # Some chart types don't support series styling
            pass
        
        # Apply custom colors from chart data
        custom_colors = chart_data.get('colors', [])
        if custom_colors:
            try:
                for i, series in enumerate(chart.series):
                    if i < len(custom_colors):
                        color_hex = custom_colors[i].lstrip('#')
                        if len(color_hex) == 6:
                            r = int(color_hex[0:2], 16)
                            g = int(color_hex[2:4], 16)
                            b = int(color_hex[4:6], 16)
                            series.format.fill.solid()
                            series.format.fill.fore_color.rgb = RGBColor(r, g, b)
            except:
                pass


class PowerPointTableGenerator:
    """Generates tables for PowerPoint slides."""
    
    @classmethod
    def add_table(
        cls,
        slide,
        table_data: Dict[str, Any],
        x: float,
        y: float,
        width: float,
        height: float
    ):
        """Add a table to the slide."""
        data = table_data.get('data', [])
        columns = table_data.get('columns', [])
        
        if not data or not columns:
            return None
        
        # Calculate dimensions
        rows = len(data) + 1  # +1 for header
        cols = len(columns)
        
        # Create table
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        table_shape = slide.shapes.add_table(rows, cols, abs_x, abs_y, abs_width, abs_height)
        table = table_shape.table
        
        # Set column headers
        for col_idx, column_name in enumerate(columns):
            cell = table.cell(0, col_idx)
            cell.text = str(column_name)
            
            # Style header
            cell.fill.solid()
            cell.fill.fore_color.rgb = PowerPointStyler.COLORS['primary']
            
            paragraph = cell.text_frame.paragraphs[0]
            font = paragraph.font
            font.color.rgb = RGBColor(255, 255, 255)  # White text
            font.bold = True
            font.size = PowerPointStyler.FONT_SIZES['body']
        
        # Fill data rows
        for row_idx, row_data in enumerate(data):
            for col_idx, cell_value in enumerate(row_data):
                if col_idx < cols:  # Ensure we don't exceed column count
                    cell = table.cell(row_idx + 1, col_idx)
                    cell.text = str(cell_value)
                    
                    # Style data cells
                    paragraph = cell.text_frame.paragraphs[0]
                    font = paragraph.font
                    font.size = PowerPointStyler.FONT_SIZES['body']
                    font.color.rgb = PowerPointStyler.COLORS['text']
                    
                    # Alternate row colors
                    if row_idx % 2 == 1:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = PowerPointStyler.COLORS['light_gray']
        
        return table


class PowerPointImageHandler:
    """Handles image insertion in PowerPoint slides."""
    
    @classmethod
    def add_image(
        cls,
        slide,
        image_path: str,
        x: float,
        y: float,
        width: float,
        height: float,
        fallback_path: str = None
    ):
        """Add an image to the slide."""
        # Check if image exists
        if not os.path.exists(image_path):
            if fallback_path and os.path.exists(fallback_path):
                image_path = fallback_path
            else:
                # Create a placeholder rectangle
                return cls._add_image_placeholder(slide, x, y, width, height)
        
        try:
            # Add image
            abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
            
            picture = slide.shapes.add_picture(image_path, abs_x, abs_y, abs_width, abs_height)
            return picture
            
        except Exception:
            # If image loading fails, add placeholder
            return cls._add_image_placeholder(slide, x, y, width, height)
    
    @classmethod
    def _add_image_placeholder(cls, slide, x: float, y: float, width: float, height: float):
        """Add an image placeholder rectangle."""
        abs_x, abs_y, abs_width, abs_height = PowerPointLayoutManager.get_position(x, y, width, height)
        
        # Add rectangle shape
        shape = slide.shapes.add_shape(
            1,  # Rectangle
            abs_x, abs_y, abs_width, abs_height
        )
        
        # Style as placeholder
        shape.fill.solid()
        shape.fill.fore_color.rgb = PowerPointStyler.COLORS['light_gray']
        shape.line.color.rgb = PowerPointStyler.COLORS['dark_gray']
        
        # Add placeholder text
        shape.text = "Image Placeholder"
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        font = paragraph.font
        font.size = PowerPointStyler.FONT_SIZES['caption']
        font.color.rgb = PowerPointStyler.COLORS['dark_gray']
        
        return shape